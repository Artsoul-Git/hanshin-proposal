#!/usr/bin/env python3
"""
動的ガントチャート v2.0 - 大吉製作所 経理部門 業務改善・AI活用 伴走支援
===
【動的機能】
- 開始日・終了日の変更で日数が自動計算
- 日付変更でガントバーが自動で伸縮
- ステータス変更でバーの色が自動変更（条件付き書式）

【Jony Iveデザインシステム準拠】
- 未着手: #CCCCCC（ペールグレー）
- 進行中: #666666（ミディアムグレー）
- 完了: #333333（ダークグレー）
- 保留: #FFFFFF（ホワイト＋枠線）
- マイルストーン: #5B7B94（アクセントブルー）
"""

from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from openpyxl.utils import get_column_letter
from openpyxl.formatting.rule import FormulaRule, CellIsRule
from openpyxl.worksheet.datavalidation import DataValidation
from datetime import datetime, timedelta

# Jony Ive Design System Colors
COLORS = {
    'black': '000000',
    'dark_gray': '333333',
    'medium_gray': '666666',
    'light_gray': '999999',
    'pale_gray': 'CCCCCC',
    'separator': 'E0E0E0',
    'bg_gray': 'F5F5F5',
    'white': 'FFFFFF',
    'accent': '5B7B94',
    # ステータス
    'status_pending': 'CCCCCC',    # 未着手
    'status_progress': '666666',   # 進行中
    'status_done': '333333',       # 完了
    'status_hold': 'FFFFFF',       # 保留
    'milestone': '5B7B94',
}

PROJECT_DATA = {
    "project_name": "大吉製作所 経理部門 業務改善・AI活用 伴走支援",
    "start_date": datetime(2025, 1, 6),
    "end_date": datetime(2025, 6, 30),
    "responsible": "株式会社けんけん",
    "client": "大吉製作所（大吉合同会社）",
    
    "phases": [
        {
            "name": "Phase 1: 現状分析",
            "tasks": [
                {"wbs": "1.1", "name": "キックオフミーティング", "assignee": "全員", "days": 1},
                {"wbs": "1.2", "name": "経理業務の棚卸し", "assignee": "乙・MIRAI", "days": 10},
                {"wbs": "1.3", "name": "業務フロー可視化", "assignee": "乙", "days": 10},
                {"wbs": "1.4", "name": "課題分析レポート作成", "assignee": "乙", "days": 5},
            ],
            "milestone": "現状分析完了",
        },
        {
            "name": "Phase 2: 型の設計",
            "tasks": [
                {"wbs": "2.1", "name": "原価集計「型」設計", "assignee": "乙・佐藤", "days": 10},
                {"wbs": "2.2", "name": "月次締め「型」設計", "assignee": "乙・MIRAI", "days": 10},
                {"wbs": "2.3", "name": "型の検証・調整", "assignee": "全員", "days": 5},
            ],
            "milestone": "型の設計完了",
        },
        {
            "name": "Phase 3: AI導入PoC",
            "tasks": [
                {"wbs": "3.1", "name": "AI下書き機能のPoC設計", "assignee": "乙", "days": 5},
                {"wbs": "3.2", "name": "PoCツール選定・設定", "assignee": "乙", "days": 5},
                {"wbs": "3.3", "name": "PoC実施（原価集計）", "assignee": "乙・佐藤", "days": 10},
                {"wbs": "3.4", "name": "PoC結果分析", "assignee": "乙", "days": 5},
            ],
            "milestone": "PoC完了",
        },
        {
            "name": "Phase 4: 実運用テスト",
            "tasks": [
                {"wbs": "4.1", "name": "運用ルール整備", "assignee": "乙・MIRAI", "days": 5},
                {"wbs": "4.2", "name": "実運用テスト（月次1回目）", "assignee": "経理チーム", "days": 10},
                {"wbs": "4.3", "name": "改善点洗い出し", "assignee": "全員", "days": 5},
                {"wbs": "4.4", "name": "ルール修正", "assignee": "乙", "days": 5},
            ],
            "milestone": "実運用テスト完了",
        },
        {
            "name": "Phase 5: 引き継ぎ設計",
            "tasks": [
                {"wbs": "5.1", "name": "属人業務の切り出し", "assignee": "乙・MIRAI", "days": 10},
                {"wbs": "5.2", "name": "引き継ぎマニュアル作成", "assignee": "乙", "days": 10},
                {"wbs": "5.3", "name": "AI活用研修実施", "assignee": "乙", "days": 5},
            ],
            "milestone": "引き継ぎ準備完了",
        },
        {
            "name": "Phase 6: 定着・展開",
            "tasks": [
                {"wbs": "6.1", "name": "実運用（月次2回目）", "assignee": "経理チーム", "days": 10},
                {"wbs": "6.2", "name": "最終調整", "assignee": "乙", "days": 5},
                {"wbs": "6.3", "name": "成果報告・完了報告", "assignee": "乙", "days": 3},
            ],
            "milestone": "プロジェクト完了",
        },
    ],
}

def create_dynamic_gantt():
    wb = Workbook()
    
    # シート1: プロジェクト概要
    ws1 = wb.active
    ws1.title = "概要"
    ws1.sheet_view.showGridLines = False
    
    title_font = Font(name='メイリオ', size=16, bold=True, color=COLORS['black'])
    heading_font = Font(name='メイリオ', size=11, bold=True, color=COLORS['dark_gray'])
    body_font = Font(name='メイリオ', size=10, color=COLORS['dark_gray'])
    small_font = Font(name='メイリオ', size=9, color=COLORS['light_gray'])
    
    ws1.column_dimensions['A'].width = 5
    ws1.column_dimensions['B'].width = 20
    ws1.column_dimensions['C'].width = 50
    
    row = 2
    ws1.merge_cells(f'B{row}:C{row}')
    ws1[f'B{row}'] = PROJECT_DATA['project_name']
    ws1[f'B{row}'].font = title_font
    ws1.row_dimensions[row].height = 35
    
    row += 2
    info = [
        ('クライアント', PROJECT_DATA['client']),
        ('担当', PROJECT_DATA['responsible']),
        ('期間', f"{PROJECT_DATA['start_date'].strftime('%Y年%m月%d日')} 〜 {PROJECT_DATA['end_date'].strftime('%Y年%m月%d日')}"),
        ('総日数', f"{(PROJECT_DATA['end_date'] - PROJECT_DATA['start_date']).days + 1}日（約{(PROJECT_DATA['end_date'] - PROJECT_DATA['start_date']).days // 7}週）"),
    ]
    
    for label, value in info:
        ws1[f'B{row}'] = label
        ws1[f'B{row}'].font = heading_font
        ws1[f'C{row}'] = value
        ws1[f'C{row}'].font = body_font
        row += 1
    
    row += 1
    ws1[f'B{row}'] = '【動的機能】'
    ws1[f'B{row}'].font = heading_font
    row += 1
    
    features = [
        '・開始日/終了日を変更 → 日数が自動計算',
        '・日付変更 → ガントバーが自動で伸縮',
        '・ステータス変更 → バーの色が自動変更',
    ]
    for feature in features:
        ws1[f'B{row}'] = feature
        ws1[f'B{row}'].font = body_font
        ws1.merge_cells(f'B{row}:C{row}')
        row += 1
    
    row += 1
    ws1[f'B{row}'] = '【ステータス凡例】'
    ws1[f'B{row}'].font = heading_font
    row += 1
    
    statuses = [
        ('未着手', COLORS['status_pending'], '存在するが目立たない'),
        ('進行中', COLORS['status_progress'], '注目すべき（濃いグレー）'),
        ('完了', COLORS['status_done'], '確定・安定（最も濃い）'),
        ('保留', COLORS['status_hold'], '空の状態（白＋枠線）'),
    ]
    for status, color, desc in statuses:
        ws1[f'B{row}'] = f'■ {status}'
        ws1[f'B{row}'].font = Font(name='メイリオ', size=10, color=color if color != 'FFFFFF' else '333333')
        ws1[f'C{row}'] = desc
        ws1[f'C{row}'].font = small_font
        row += 1
    
    # シート2: 動的ガントチャート
    ws2 = wb.create_sheet("WBS・ガントチャート")
    ws2.sheet_view.showGridLines = False
    
    # 列幅設定
    ws2.column_dimensions['A'].width = 6   # WBS
    ws2.column_dimensions['B'].width = 30  # タスク名
    ws2.column_dimensions['C'].width = 12  # 担当
    ws2.column_dimensions['D'].width = 12  # 開始日
    ws2.column_dimensions['E'].width = 12  # 終了日
    ws2.column_dimensions['F'].width = 6   # 日数（関数）
    ws2.column_dimensions['G'].width = 10  # ステータス
    
    # 週列（26週間）
    START_COL = 8  # H列から
    WEEKS = 26
    for i in range(WEEKS):
        ws2.column_dimensions[get_column_letter(START_COL + i)].width = 4
    
    # プロジェクト開始日をセル参照用に記録（I1に隠し保存）
    ws2['I1'] = PROJECT_DATA['start_date']
    ws2['I1'].number_format = 'YYYY/MM/DD'
    ws2['I1'].font = Font(color='FFFFFF')  # 白文字で非表示
    
    # ヘッダー行
    row = 2
    headers = ['WBS', 'タスク名', '担当', '開始日', '終了日', '日数', 'ステータス']
    for col, header in enumerate(headers, 1):
        cell = ws2.cell(row=row, column=col, value=header)
        cell.font = heading_font
        cell.fill = PatternFill(start_color=COLORS['bg_gray'], end_color=COLORS['bg_gray'], fill_type='solid')
        cell.alignment = Alignment(horizontal='center', vertical='center')
        cell.border = Border(bottom=Side(style='thin', color=COLORS['separator']))
    
    # 週ヘッダー
    for i in range(WEEKS):
        col = START_COL + i
        week_label = f"W{i+1}"
        cell = ws2.cell(row=row, column=col, value=week_label)
        cell.font = small_font
        cell.fill = PatternFill(start_color=COLORS['bg_gray'], end_color=COLORS['bg_gray'], fill_type='solid')
        cell.alignment = Alignment(horizontal='center', vertical='center')
        cell.border = Border(bottom=Side(style='thin', color=COLORS['separator']))
    
    ws2.row_dimensions[row].height = 25
    
    # ステータスのドロップダウン
    status_dv = DataValidation(
        type="list",
        formula1='"未着手,進行中,完了,保留"',
        allow_blank=True
    )
    ws2.add_data_validation(status_dv)
    
    # 薄いボーダー
    thin_border = Border(
        left=Side(style='thin', color=COLORS['separator']),
        right=Side(style='thin', color=COLORS['separator']),
        top=Side(style='thin', color=COLORS['separator']),
        bottom=Side(style='thin', color=COLORS['separator'])
    )
    
    # タスクデータ生成
    row = 3
    current_date = PROJECT_DATA['start_date']
    task_rows = []  # 条件付き書式適用用
    
    # タスク日数のスケーリング計算
    total_task_days = sum(
        sum(task['days'] for task in phase['tasks'])
        for phase in PROJECT_DATA['phases']
    )
    total_project_days = (PROJECT_DATA['end_date'] - PROJECT_DATA['start_date']).days + 1
    scale_factor = total_project_days / total_task_days * 0.85
    
    for phase_idx, phase in enumerate(PROJECT_DATA['phases']):
        # フェーズ行
        ws2.merge_cells(f'A{row}:G{row}')
        phase_cell = ws2.cell(row=row, column=1, value=phase['name'])
        phase_cell.font = Font(name='メイリオ', size=11, bold=True, color=COLORS['accent'])
        phase_cell.fill = PatternFill(start_color='F0F4F8', end_color='F0F4F8', fill_type='solid')
        ws2.row_dimensions[row].height = 25
        row += 1
        
        for task in phase['tasks']:
            adjusted_days = max(1, int(task['days'] * scale_factor))
            end_date = current_date + timedelta(days=adjusted_days - 1)
            
            # WBS
            ws2.cell(row=row, column=1, value=task['wbs']).font = body_font
            
            # タスク名
            ws2.cell(row=row, column=2, value=task['name']).font = body_font
            
            # 担当
            ws2.cell(row=row, column=3, value=task['assignee']).font = body_font
            
            # 開始日（編集可能）
            start_cell = ws2.cell(row=row, column=4, value=current_date)
            start_cell.font = body_font
            start_cell.number_format = 'MM/DD'
            
            # 終了日（編集可能）
            end_cell = ws2.cell(row=row, column=5, value=end_date)
            end_cell.font = body_font
            end_cell.number_format = 'MM/DD'
            
            # 日数（関数で自動計算）
            days_cell = ws2.cell(row=row, column=6)
            days_cell.value = f'=E{row}-D{row}+1'
            days_cell.font = body_font
            days_cell.alignment = Alignment(horizontal='center')
            
            # ステータス
            status_cell = ws2.cell(row=row, column=7, value='未着手')
            status_cell.font = body_font
            status_dv.add(status_cell)
            
            # ガントバー（関数で動的に表示）
            for w in range(WEEKS):
                col = START_COL + w
                week_start = PROJECT_DATA['start_date'] + timedelta(days=w * 7)
                week_end = week_start + timedelta(days=6)
                
                bar_cell = ws2.cell(row=row, column=col, value='')
                
                # 関数でバー表示を制御
                # IF(AND(開始日<=週末, 終了日>=週開始), "■", "")
                week_start_str = week_start.strftime('%Y/%m/%d')
                week_end_str = week_end.strftime('%Y/%m/%d')
                
                formula = f'=IF(AND($D{row}<=DATE({week_end.year},{week_end.month},{week_end.day}),$E{row}>=DATE({week_start.year},{week_start.month},{week_start.day})),"■","")'
                bar_cell.value = formula
                bar_cell.font = Font(name='メイリオ', size=10, color=COLORS['pale_gray'])
                bar_cell.alignment = Alignment(horizontal='center', vertical='center')
                bar_cell.border = thin_border
            
            task_rows.append(row)
            ws2.row_dimensions[row].height = 22
            row += 1
            
            current_date = end_date + timedelta(days=1)
        
        # マイルストーン行
        milestone_date = current_date - timedelta(days=1)
        
        ws2.cell(row=row, column=1, value='').font = body_font
        ms_name_cell = ws2.cell(row=row, column=2, value=f"◆ {phase['milestone']}")
        ms_name_cell.font = Font(name='メイリオ', size=10, bold=True, color=COLORS['milestone'])
        
        ms_date_cell = ws2.cell(row=row, column=4, value=milestone_date)
        ms_date_cell.font = body_font
        ms_date_cell.number_format = 'MM/DD'
        
        # マイルストーンのガントバー
        milestone_week = (milestone_date - PROJECT_DATA['start_date']).days // 7
        if 0 <= milestone_week < WEEKS:
            ms_bar = ws2.cell(row=row, column=START_COL + milestone_week, value='◆')
            ms_bar.font = Font(name='メイリオ', size=12, bold=True, color='FFFFFF')
            ms_bar.fill = PatternFill(start_color=COLORS['milestone'], end_color=COLORS['milestone'], fill_type='solid')
            ms_bar.alignment = Alignment(horizontal='center', vertical='center')
        
        ws2.row_dimensions[row].height = 22
        row += 1
        
        # フェーズ間の調整期間
        current_date = current_date + timedelta(days=3)
    
    # 条件付き書式を適用（ステータスに応じてバーの色を変更）
    # 各タスク行のガント列に条件付き書式を設定
    
    for task_row in task_rows:
        # ガント列の範囲（H列〜AG列 = W1〜W26）
        gantt_range = f'{get_column_letter(START_COL)}{task_row}:{get_column_letter(START_COL + WEEKS - 1)}{task_row}'
        status_ref = f'$G${task_row}'
        
        # 未着手: ペールグレー (#CCCCCC)
        ws2.conditional_formatting.add(
            gantt_range,
            FormulaRule(
                formula=[f'AND({status_ref}="未着手", {get_column_letter(START_COL)}{task_row}<>"")'],
                fill=PatternFill(start_color=COLORS['status_pending'], end_color=COLORS['status_pending'], fill_type='solid'),
                font=Font(color=COLORS['status_pending'])
            )
        )
        
        # 進行中: ミディアムグレー (#666666)
        ws2.conditional_formatting.add(
            gantt_range,
            FormulaRule(
                formula=[f'AND({status_ref}="進行中", {get_column_letter(START_COL)}{task_row}<>"")'],
                fill=PatternFill(start_color=COLORS['status_progress'], end_color=COLORS['status_progress'], fill_type='solid'),
                font=Font(color=COLORS['status_progress'])
            )
        )
        
        # 完了: ダークグレー (#333333)
        ws2.conditional_formatting.add(
            gantt_range,
            FormulaRule(
                formula=[f'AND({status_ref}="完了", {get_column_letter(START_COL)}{task_row}<>"")'],
                fill=PatternFill(start_color=COLORS['status_done'], end_color=COLORS['status_done'], fill_type='solid'),
                font=Font(color=COLORS['status_done'])
            )
        )
        
        # 保留: ホワイト + 枠線
        ws2.conditional_formatting.add(
            gantt_range,
            FormulaRule(
                formula=[f'AND({status_ref}="保留", {get_column_letter(START_COL)}{task_row}<>"")'],
                fill=PatternFill(start_color=COLORS['status_hold'], end_color=COLORS['status_hold'], fill_type='solid'),
                font=Font(color=COLORS['light_gray']),
                border=Border(
                    left=Side(style='thin', color=COLORS['light_gray']),
                    right=Side(style='thin', color=COLORS['light_gray']),
                    top=Side(style='thin', color=COLORS['light_gray']),
                    bottom=Side(style='thin', color=COLORS['light_gray'])
                )
            )
        )
    
    # 凡例シート
    ws3 = wb.create_sheet("凡例・使い方")
    ws3.sheet_view.showGridLines = False
    
    ws3.column_dimensions['B'].width = 15
    ws3.column_dimensions['C'].width = 50
    
    row = 2
    ws3[f'B{row}'] = '動的ガントチャート 使い方'
    ws3[f'B{row}'].font = title_font
    ws3.merge_cells(f'B{row}:C{row}')
    
    row += 2
    ws3[f'B{row}'] = '【操作方法】'
    ws3[f'B{row}'].font = heading_font
    
    row += 1
    instructions = [
        ('開始日を変更', 'D列の日付を編集 → 日数とバーが自動更新'),
        ('終了日を変更', 'E列の日付を編集 → 日数とバーが自動更新'),
        ('ステータス変更', 'G列のドロップダウンから選択 → バー色が自動変更'),
    ]
    for action, result in instructions:
        ws3[f'B{row}'] = action
        ws3[f'B{row}'].font = body_font
        ws3[f'C{row}'] = result
        ws3[f'C{row}'].font = small_font
        row += 1
    
    row += 1
    ws3[f'B{row}'] = '【色の意味】'
    ws3[f'B{row}'].font = heading_font
    
    row += 1
    for status, color, desc in statuses:
        ws3.cell(row=row, column=2, value=status).font = body_font
        sample_cell = ws3.cell(row=row, column=3, value='■■■■■')
        sample_cell.font = Font(name='メイリオ', size=10, color=color if color != 'FFFFFF' else COLORS['light_gray'])
        if color != 'FFFFFF':
            sample_cell.fill = PatternFill(start_color=color, end_color=color, fill_type='solid')
        row += 1
    
    # 保存
    output_path = '/mnt/user-data/outputs/gantt_dynamic_大吉製作所_20251215.xlsx'
    wb.save(output_path)
    print(f"✅ 動的ガントチャート作成完了: {output_path}")
    return output_path

if __name__ == "__main__":
    create_dynamic_gantt()
