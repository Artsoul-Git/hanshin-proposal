#!/usr/bin/env python3
"""
hearing_sheet_extractor.py - ヒアリングシートからデータを抽出

hearing-sheet-generatorで生成されたExcelファイルから
proposal-outline-generator用のデータを自動抽出する。

使用例:
    from hearing_sheet_extractor import extract_from_hearing_sheet
    
    data = extract_from_hearing_sheet("ヒアリングシート_株式会社サンプル.xlsx")
    print(data["client_name"])  # 株式会社サンプル
    print(data["issues"])       # [{"title": "...", "detail": "..."}, ...]
"""

import re
from pathlib import Path
from typing import Dict, List, Any, Optional

import openpyxl
from openpyxl.utils import get_column_letter


# =============================================================================
# hearing-sheetのセル位置マッピング
# =============================================================================

# 基本情報シートのマッピング
BASIC_INFO_MAP = {
    "company_name": "B3",      # 会社名
    "location": "B4",          # 所在地
    "established": "B5",       # 設立年
    "employees": "D5",         # 従業員数
    "business": "B6",          # 事業内容
    "revenue": "B7",           # 年商規模
    "contact_dept": "B10",     # 担当者部署
    "contact_name": "B11",     # 担当者氏名
    "contact_info": "B12",     # 連絡先
}

# 優先度・ゴール整理シートのマッピング
PRIORITY_MAP = {
    "priority_1": "B3",        # 課題優先順位1
    "priority_2": "B4",        # 課題優先順位2
    "priority_3": "B5",        # 課題優先順位3
    "ideal_state": "B8",       # 理想の状態
    "budget": "B11",           # 予算感
    "schedule": "B14",         # 希望スケジュール
    "decision_process": "B17", # 意思決定プロセス
}

# 部門別シートの「困っていること」チェックボックス行範囲
DEPARTMENT_ISSUES_RANGE = {
    "経理": {"start_row": 15, "end_row": 23, "col": "A"},
    "人事・総務": {"start_row": 15, "end_row": 23, "col": "A"},
    "営業": {"start_row": 15, "end_row": 22, "col": "A"},
    "現場・製造": {"start_row": 15, "end_row": 22, "col": "A"},
    "情報システム": {"start_row": 15, "end_row": 22, "col": "A"},
}

# 部門別シートの「具体的なエピソード」セル
DEPARTMENT_EPISODE_CELL = "B25"


# =============================================================================
# 抽出関数
# =============================================================================

def extract_from_hearing_sheet(xlsx_path: str) -> Dict[str, Any]:
    """
    ヒアリングシートからデータを抽出
    
    Args:
        xlsx_path: ヒアリングシートのパス
    
    Returns:
        {
            "client_name": str,
            "industry": str,
            "employees": str,
            "issues": [{"title": str, "detail": str}, ...],
            "ideal_state": str,
            "budget": str,
            "schedule": str,
            "episodes": [str, ...],
            "raw_data": dict  # 全抽出データ
        }
    """
    wb = openpyxl.load_workbook(xlsx_path, data_only=True)
    
    result = {
        "client_name": "",
        "industry": "",
        "employees": "",
        "issues": [],
        "ideal_state": "",
        "budget": "",
        "schedule": "",
        "episodes": [],
        "raw_data": {}
    }
    
    # 基本情報シートから抽出
    if "基本情報" in wb.sheetnames:
        ws = wb["基本情報"]
        result["client_name"] = _get_cell_value(ws, BASIC_INFO_MAP["company_name"])
        result["employees"] = _get_cell_value(ws, BASIC_INFO_MAP["employees"])
        result["raw_data"]["basic_info"] = {
            key: _get_cell_value(ws, cell) 
            for key, cell in BASIC_INFO_MAP.items()
        }
        
        # 業種はシート名から推測（業種別質問がある場合）
        result["industry"] = _detect_industry_from_sheets(wb)
    
    # 優先度・ゴール整理シートから抽出
    priority_sheet_names = ["優先度・ゴール整理", "優先度整理", "ゴール整理"]
    for sheet_name in priority_sheet_names:
        if sheet_name in wb.sheetnames:
            ws = wb[sheet_name]
            
            # 課題を抽出
            priorities = []
            for key in ["priority_1", "priority_2", "priority_3"]:
                value = _get_cell_value(ws, PRIORITY_MAP[key])
                if value:
                    priorities.append({"title": value, "detail": ""})
            
            if priorities:
                result["issues"] = priorities
            
            result["ideal_state"] = _get_cell_value(ws, PRIORITY_MAP["ideal_state"])
            result["budget"] = _get_cell_value(ws, PRIORITY_MAP["budget"])
            result["schedule"] = _get_cell_value(ws, PRIORITY_MAP["schedule"])
            
            result["raw_data"]["priority"] = {
                key: _get_cell_value(ws, cell) 
                for key, cell in PRIORITY_MAP.items()
            }
            break
    
    # 部門別シートから課題とエピソードを抽出
    department_issues = []
    episodes = []
    
    for dept_name, config in DEPARTMENT_ISSUES_RANGE.items():
        if dept_name in wb.sheetnames:
            ws = wb[dept_name]
            
            # チェックされた課題を抽出
            issues = _extract_checked_issues(ws, config)
            department_issues.extend(issues)
            
            # エピソードを抽出
            episode = _get_cell_value(ws, DEPARTMENT_EPISODE_CELL)
            if episode:
                episodes.append(f"【{dept_name}】{episode}")
    
    # 優先度シートの課題がなければ、部門別の課題を使用
    if not result["issues"] and department_issues:
        result["issues"] = department_issues[:5]
    elif department_issues:
        # 部門別課題を詳細として追加
        for i, dept_issue in enumerate(department_issues):
            if i < len(result["issues"]):
                if not result["issues"][i]["detail"]:
                    result["issues"][i]["detail"] = dept_issue["title"]
    
    result["episodes"] = episodes
    result["raw_data"]["department_issues"] = department_issues
    
    wb.close()
    
    return result


def _get_cell_value(ws, cell_ref: str) -> str:
    """セルの値を文字列として取得"""
    try:
        value = ws[cell_ref].value
        if value is None:
            return ""
        return str(value).strip()
    except Exception:
        return ""


def _detect_industry_from_sheets(wb) -> str:
    """シート名から業種を推測"""
    industry_keywords = {
        "製造": "製造業",
        "小売": "小売業",
        "卸売": "卸売業",
        "建設": "建設業",
        "福祉": "福祉・介護",
        "介護": "福祉・介護",
        "医療": "医療",
        "金融": "金融・保険",
        "保険": "金融・保険",
        "不動産": "不動産",
        "運輸": "運輸・物流",
        "物流": "運輸・物流",
        "飲食": "飲食・宿泊",
        "宿泊": "飲食・宿泊",
        "IT": "情報通信・IT",
        "情報": "情報通信・IT",
        "農": "農林水産業",
        "漁": "農林水産業",
        "公共": "公共・団体",
        "農協": "公共・団体",
        "漁協": "公共・団体",
        "商工会": "公共・団体",
    }
    
    # 基本情報シートの事業内容からも推測
    if "基本情報" in wb.sheetnames:
        ws = wb["基本情報"]
        business = _get_cell_value(ws, BASIC_INFO_MAP["business"])
        for keyword, industry in industry_keywords.items():
            if keyword in business:
                return industry
    
    # シート名から推測
    for sheet_name in wb.sheetnames:
        for keyword, industry in industry_keywords.items():
            if keyword in sheet_name:
                return industry
    
    return "その他サービス業"


def _extract_checked_issues(ws, config: dict) -> List[Dict[str, str]]:
    """
    チェックボックス形式の課題を抽出
    
    hearing-sheetでは「□ 課題名」の形式で記載されており、
    チェック済みは「☑」や「■」に変わっている想定
    """
    issues = []
    col = config["col"]
    
    for row in range(config["start_row"], config["end_row"] + 1):
        cell_value = _get_cell_value(ws, f"{col}{row}")
        
        if not cell_value:
            continue
        
        # チェック済みマーカーを検出
        checked_markers = ["☑", "■", "✓", "✔", "○", "●", "[x]", "[X]"]
        is_checked = any(marker in cell_value for marker in checked_markers)
        
        if is_checked:
            # マーカーを除去して課題名を抽出
            issue_text = cell_value
            for marker in checked_markers + ["□", "[ ]"]:
                issue_text = issue_text.replace(marker, "")
            issue_text = issue_text.strip()
            
            if issue_text:
                issues.append({
                    "title": issue_text,
                    "detail": ""
                })
    
    return issues


def generate_proposal_data(hearing_data: Dict[str, Any]) -> Dict[str, Any]:
    """
    ヒアリングデータから提案骨子用データを生成
    
    Args:
        hearing_data: extract_from_hearing_sheetの戻り値
    
    Returns:
        create_proposal_outlineに渡すデータ
    """
    # 提案タイプを判定
    all_text = " ".join([
        hearing_data.get("ideal_state", ""),
        " ".join(hearing_data.get("episodes", [])),
        " ".join([i["title"] for i in hearing_data.get("issues", [])])
    ])
    
    proposal_type = _determine_proposal_type(all_text)
    
    # アプローチを生成
    approach = _generate_approach(hearing_data, proposal_type)
    
    # 理由を生成
    reasons = _generate_reasons(hearing_data, proposal_type)
    
    # 期待効果を生成
    effects = _generate_effects(hearing_data)
    
    # 次のステップを生成
    next_steps = _generate_next_steps(hearing_data)
    
    return {
        "client_name": hearing_data.get("client_name", ""),
        "issues": hearing_data.get("issues", []),
        "proposal_type": proposal_type,
        "approach": approach,
        "reasons": reasons,
        "effects": effects,
        "next_steps": next_steps,
    }


def _determine_proposal_type(text: str) -> str:
    """提案タイプを判定"""
    ai_keywords = ["AI", "生成AI", "ChatGPT", "自動化", "効率化", "DX", "デジタル化", "ペーパーレス", "RPA"]
    it_keywords = ["システム導入", "ツール導入", "クラウド", "ソフトウェア", "基幹システム"]
    
    text_upper = text.upper()
    ai_score = sum(1 for kw in ai_keywords if kw.upper() in text_upper)
    it_score = sum(1 for kw in it_keywords if kw.upper() in text_upper)
    
    # デフォルトはAI/DX（けんけんの主力サービス）
    if ai_score >= it_score:
        return "AI/DXアドバイザリー"
    else:
        return "IT導入支援"


def _generate_approach(hearing_data: Dict, proposal_type: str) -> str:
    """アプローチ文を生成"""
    issues = hearing_data.get("issues", [])
    ideal = hearing_data.get("ideal_state", "")
    
    if not issues:
        return "業務プロセスの可視化と改善提案"
    
    main_issue = issues[0]["title"] if issues else "業務効率化"
    
    if proposal_type == "AI/DXアドバイザリー":
        return f"生成AIを活用した{main_issue}の改善と業務フローの最適化"
    else:
        return f"ITツール導入による{main_issue}の解決と業務効率化"


def _generate_reasons(hearing_data: Dict, proposal_type: str) -> List[str]:
    """アプローチの理由を生成"""
    base_reasons = [
        "手作業の削減による工数・コスト削減効果が大きい",
        "既存システムとの連携が容易",
        "段階的な導入でリスクを最小化できる",
    ]
    
    employees = hearing_data.get("employees", "")
    if employees and any(c.isdigit() for c in employees):
        # 従業員数から規模を判断
        num = int("".join(filter(str.isdigit, employees)))
        if num <= 10:
            base_reasons[0] = "少人数でも大きな効果が得られる"
        elif num <= 50:
            base_reasons[0] = "組織規模に適した段階的な導入が可能"
    
    return base_reasons


def _generate_effects(hearing_data: Dict) -> List[Dict[str, str]]:
    """期待効果を生成"""
    issues = hearing_data.get("issues", [])
    
    effects = []
    
    # 課題に基づいて効果を推定
    issue_effect_map = {
        "手入力": {"item": "工数削減", "value": "約60%削減"},
        "時間がかかる": {"item": "処理時間", "value": "約50%短縮"},
        "属人化": {"item": "業務標準化", "value": "ナレッジ共有実現"},
        "連携": {"item": "データ連携", "value": "リアルタイム化"},
        "ミス": {"item": "品質向上", "value": "エラー率大幅削減"},
        "紙": {"item": "ペーパーレス", "value": "紙コスト削減"},
    }
    
    for issue in issues[:3]:
        title = issue.get("title", "")
        for keyword, effect in issue_effect_map.items():
            if keyword in title:
                if effect not in effects:
                    effects.append(effect)
                break
    
    # 最低3つは効果を出す
    default_effects = [
        {"item": "工数削減", "value": "約50%削減"},
        {"item": "コスト", "value": "年間削減効果あり"},
        {"item": "品質向上", "value": "ミス・手戻り削減"},
    ]
    
    while len(effects) < 3:
        for de in default_effects:
            if de not in effects:
                effects.append(de)
                break
    
    return effects[:3]


def _generate_next_steps(hearing_data: Dict) -> List[str]:
    """次のステップを生成"""
    schedule = hearing_data.get("schedule", "")
    
    steps = [
        "詳細ヒアリング（日程調整中）",
        "正式見積のご提示",
        "PoC（概念実証）のご相談",
    ]
    
    if schedule:
        steps[0] = f"詳細ヒアリング（{schedule}を目途）"
    
    return steps


# =============================================================================
# テスト・デバッグ用
# =============================================================================

if __name__ == "__main__":
    import sys
    
    if len(sys.argv) < 2:
        print("Usage: python hearing_sheet_extractor.py <hearing_sheet.xlsx>")
        sys.exit(1)
    
    xlsx_path = sys.argv[1]
    
    print(f"Extracting from: {xlsx_path}")
    print("-" * 50)
    
    data = extract_from_hearing_sheet(xlsx_path)
    
    print(f"Client: {data['client_name']}")
    print(f"Industry: {data['industry']}")
    print(f"Employees: {data['employees']}")
    print(f"Issues: {len(data['issues'])} items")
    for i, issue in enumerate(data['issues'], 1):
        print(f"  {i}. {issue['title']}")
    print(f"Ideal State: {data['ideal_state']}")
    print(f"Budget: {data['budget']}")
    print(f"Episodes: {len(data['episodes'])} items")
    
    print("-" * 50)
    print("Proposal Data:")
    proposal_data = generate_proposal_data(data)
    print(f"  Type: {proposal_data['proposal_type']}")
    print(f"  Approach: {proposal_data['approach']}")
