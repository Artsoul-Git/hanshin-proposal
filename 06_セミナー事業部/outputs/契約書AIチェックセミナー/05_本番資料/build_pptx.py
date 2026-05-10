#!/usr/bin/env python
# -*- coding: utf-8 -*-
"""
契約書AIチェックセミナー PPTX Builder v2
Reads slide-spec.json → produces properly formatted 16:9 PPTX
Design: navy #1A3C6E + orange #E85D26  (matches 06_HTML版)
"""

import json, os, re
from lxml import etree
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.oxml.ns import qn

# ─── slide dimensions ────────────────────────────────────
W = Inches(13.33)
H = Inches(7.5)

# ─── palette ─────────────────────────────────────────────
NAVY    = RGBColor(0x1A, 0x3C, 0x6E)
NAVY_D  = RGBColor(0x0D, 0x22, 0x40)
NAVY_M  = RGBColor(0x2D, 0x60, 0x96)
NAVY_L  = RGBColor(0xEE, 0xF3, 0xF9)
NAVY_20 = RGBColor(0xBD, 0xD0, 0xE6)
ORANGE  = RGBColor(0xE8, 0x5D, 0x26)
ORG_L   = RGBColor(0xFD, 0xF1, 0xEC)
WHITE   = RGBColor(0xFF, 0xFF, 0xFF)
TEXT    = RGBColor(0x1E, 0x27, 0x35)
SUB     = RGBColor(0x6B, 0x72, 0x80)
LINE    = RGBColor(0xD1, 0xD9, 0xE0)
GRAY_L  = RGBColor(0xF8, 0xF9, 0xFA)
RISK_H  = RGBColor(0xC0, 0x39, 0x2B)
RISK_M  = RGBColor(0xE6, 0x7E, 0x22)
DIM_W   = RGBColor(0xCC, 0xD9, 0xEA)
MID_W   = RGBColor(0x80, 0x9D, 0xBF)

BADGE_GRN = RGBColor(0x6F, 0x91, 0x1D)
BADGE_DRK = RGBColor(0x33, 0x33, 0x33)
BADGE_RED = RGBColor(0xC0, 0x39, 0x2B)
TEXT_YEL  = RGBColor(0xFA, 0xBE, 0x00)
BG_PALE   = RGBColor(0xE8, 0xF1, 0xD8)

FONT = 'Meiryo'

# ─── layout constants ────────────────────────────────────
ML = Inches(0.85)      # margin left
MR = Inches(0.85)      # margin right
MT = Inches(0.45)      # margin top
CW = W - ML - MR       # content width
HDR_H = Inches(0.95)   # header block height
HDR_BAR = Pt(3)        # orange underline bar thickness


# ─── xml helpers ─────────────────────────────────────────

def _no_border(shape):
    """Remove all borders from a shape via XML."""
    sp_pr = shape._element.spPr
    for ln in sp_pr.findall(qn('a:ln')):
        sp_pr.remove(ln)
    ln = etree.SubElement(sp_pr, qn('a:ln'))
    etree.SubElement(ln, qn('a:noFill'))


def _set_para_spacing(para, before_pt=0, after_pt=0, line_spacing_pt=None):
    pPr = para._p.get_or_add_pPr()
    if before_pt:
        pPr.set('spcBef', str(int(before_pt * 12700)))  # 1pt = 12700 EMU
    if after_pt:
        pPr.set('spcAft', str(int(after_pt * 12700)))


# ─── drawing primitives ──────────────────────────────────

def rect(slide, x, y, w, h, fill=None, line=None, line_w_pt=0.75):
    shape = slide.shapes.add_shape(1, x, y, w, h)
    if fill:
        shape.fill.solid()
        shape.fill.fore_color.rgb = fill
    else:
        shape.fill.background()
    if line:
        shape.line.color.rgb = line
        shape.line.width = Pt(line_w_pt)
    else:
        _no_border(shape)
    return shape


def txbox(slide, x, y, w, h, text, size=16, bold=False, color=TEXT,
          align=PP_ALIGN.LEFT, font=FONT, wrap=True, italic=False):
    box = slide.shapes.add_textbox(x, y, w, h)
    tf = box.text_frame
    tf.word_wrap = wrap
    para = tf.paragraphs[0]
    para.alignment = align
    run = para.add_run()
    run.text = text
    run.font.name = font
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.italic = italic
    run.font.color.rgb = color
    return box


def set_bg(slide, color):
    bg = slide.background
    fill = bg.fill
    fill.solid()
    fill.fore_color.rgb = color


def add_notes(slide, text):
    if text and text.strip():
        slide.notes_slide.notes_text_frame.text = text


# ─── shared components ───────────────────────────────────

def header(slide, title):
    """Standard slide header: title + orange underline bar. Returns content Y."""
    txbox(slide, ML, MT, CW, Pt(30), title, size=22, bold=True, color=NAVY)
    bar_y = MT + Pt(32)
    rect(slide, ML, bar_y, CW, HDR_BAR, fill=ORANGE)
    return bar_y + HDR_BAR + Inches(0.18)


def bullets(slide, lines_data, start_y, avail_h, max_items=7):
    """
    lines_data: list of (text, level, is_callout)
      level 0 = main bullet  (orange dot, navy bold)
      level 1 = sub item     (dash, gray text)
      level 2 = callout box  (navy left border box)
    """
    items = lines_data[:max_items]
    if not items:
        return

    # Estimate heights
    sizes = {0: 16, 1: 13, 2: 14, 3: 16, 4: 16, 5: 16}
    total_pts = sum(sizes.get(lvl, 14) * 1.6 + 3 for _, lvl, _ in items)
    available_pts = avail_h / Pt(1)
    scale = min(1.0, available_pts / total_pts) if total_pts > 0 else 1.0

    y = start_y
    for text, level, is_callout in items:
        base_size = sizes.get(level, 14)
        sz = max(10, int(base_size * scale))
        item_h = Pt(sz * 1.65)

        if is_callout or level == 2:
            # Callout box with left border
            rect(slide, ML, y, CW, item_h + Pt(4), fill=NAVY_L)
            rect(slide, ML, y, Pt(4), item_h + Pt(4), fill=NAVY)
            txbox(slide, ML + Pt(10), y + Pt(2), CW - Pt(10), item_h,
                  text, size=sz, bold=True, color=NAVY)
        elif level == 3: # ケース
            badge_text, content_text = text
            badge_w = Pt(90)
            rect(slide, ML, y, badge_w, item_h + Pt(4), fill=BADGE_GRN)
            txbox(slide, ML, y + Pt(2), badge_w, item_h, badge_text, size=sz, bold=True, color=WHITE, align=PP_ALIGN.CENTER)
            content_w = CW - badge_w - Pt(10)
            rect(slide, ML + badge_w + Pt(10), y, content_w, item_h + Pt(4), fill=BADGE_RED)
            txbox(slide, ML + badge_w + Pt(15), y + Pt(2), content_w - Pt(5), item_h, "⚠ " + content_text, size=sz, bold=True, color=WHITE)
        elif level == 4: # ルール
            badge_text, content_text = text
            badge_w = Pt(90)
            badge_color = BADGE_DRK if '②' in badge_text else BADGE_GRN
            text_color = TEXT_YEL if '②' in badge_text else NAVY
            rect(slide, ML, y, badge_w, item_h + Pt(4), fill=badge_color)
            txbox(slide, ML, y + Pt(2), badge_w, item_h, badge_text, size=sz, bold=True, color=WHITE, align=PP_ALIGN.CENTER)
            txbox(slide, ML + badge_w + Pt(10), y + Pt(2), CW - badge_w - Pt(10), item_h, content_text, size=sz, bold=True, color=text_color)
        elif level == 5: # 題材バナー
            topic, tool = text
            banner_h = item_h + Pt(10)
            rect(slide, ML, y, CW, banner_h, fill=BG_PALE)
            txbox(slide, ML + Pt(10), y + Pt(5), CW * 0.65, item_h, "★ 題材：" + topic, size=sz, bold=True, color=BADGE_GRN)
            if tool:
                txbox(slide, ML + CW * 0.65, y + Pt(5), CW * 0.35 - Pt(10), item_h, "｜ ツール：" + tool, size=sz, bold=True, color=BADGE_GRN)
            y += Pt(10) # extra height for banner
        elif level == 0:
            txbox(slide, ML, y, Pt(10), item_h, '●', size=sz - 2, color=ORANGE)
            txbox(slide, ML + Pt(14), y, CW - Pt(14), item_h,
                  text, size=sz, bold=True, color=NAVY)
        else:
            txbox(slide, ML + Pt(8), y, Pt(10), item_h, '─', size=sz - 2, color=LINE)
            txbox(slide, ML + Pt(22), y, CW - Pt(22), item_h,
                  text, size=sz, color=SUB)
        y += item_h + Pt(3)


def parse_body(body_text):
    """
    Parse body text into list of (text, level, is_callout).
    Lines starting with 　 (full-width space) or 2+ spaces → level 1
    Lines starting with ▶ → callout
    ケースX, ルールX, 題材： → level 3, 4, 5
    Otherwise → level 0
    """
    result = []
    import re
    for raw in body_text.split('\n'):
        if not raw.strip():
            continue
        stripped = raw.strip()
        if stripped.startswith('▶'):
            result.append((stripped[1:].strip(), 2, True))
        elif raw.startswith('　') or raw.startswith('  '):
            result.append((stripped.lstrip('→ ').strip(), 1, False))
        elif stripped.startswith('ケース'):
            m = re.match(r'^(ケース\d+)[:：]?\s*(.*)', stripped)
            if m:
                result.append(((m.group(1), m.group(2)), 3, False))
            else:
                result.append((stripped, 0, False))
        elif stripped.startswith('ルール'):
            m = re.match(r'^(ルール\d+)[:：]?\s*(.*)', stripped)
            if m:
                result.append(((m.group(1), m.group(2)), 4, False))
            else:
                result.append((stripped, 0, False))
        elif stripped.startswith('題材：'):
            m = re.match(r'^題材：\s*(.*)', stripped)
            if m:
                result.append(((m.group(1).strip(), ""), 5, False))
            else:
                result.append((stripped, 0, False))
        elif stripped.startswith('ツール：'):
            m = re.match(r'^ツール：\s*(.*)', stripped)
            if m and result and result[-1][1] == 5:
                prev_topic = result[-1][0][0]
                result[-1] = ((prev_topic, m.group(1).strip()), 5, False)
            else:
                result.append((stripped, 0, False))
        else:
            result.append((stripped, 0, False))
    return result


# ─── slide type builders ─────────────────────────────────

def slide_cover(prs, s):
    sl = prs.slides.add_slide(prs.slide_layouts[6])
    set_bg(sl, NAVY_D)

    # Top thin orange bar
    rect(sl, 0, 0, W, Pt(5), fill=ORANGE)

    # Company tag
    txbox(sl, ML, Inches(0.65), CW, Pt(18),
          'AS株式会社 セミナー事業部', size=11, color=MID_W)

    # Orange accent bar
    rect(sl, ML, Inches(2.7), Inches(1.4), Pt(4), fill=ORANGE)

    # Main title
    title = s.get('タイトル', '')
    txbox(sl, ML, Inches(3.0), CW, Inches(1.6),
          title, size=42, bold=True, color=WHITE)

    # Subtitle
    sub = s.get('サブタイトル', '')
    txbox(sl, ML, Inches(4.7), CW, Inches(0.9),
          sub, size=19, color=DIM_W)

    # Date
    date = s.get('日付', '')
    txbox(sl, ML, Inches(6.8), CW, Pt(18),
          date, size=12, color=MID_W)

    # Right accent bar
    rect(sl, W - Pt(6), 0, Pt(6), H, fill=ORANGE)

    add_notes(sl, s.get('スピーカーノート', ''))
    return sl


def slide_content(prs, s):
    sl = prs.slides.add_slide(prs.slide_layouts[6])
    set_bg(sl, WHITE)

    # Navy top bar (thin)
    rect(sl, 0, 0, W, Pt(4), fill=NAVY)

    content_y = header(sl, s.get('タイトル', ''))
    avail_h = H - content_y - Inches(0.25)

    body = s.get('本文', '')
    items = parse_body(body)
    bullets(sl, items, content_y, avail_h)

    add_notes(sl, s.get('スピーカーノート', ''))
    return sl


def slide_section(prs, s):
    sl = prs.slides.add_slide(prs.slide_layouts[6])
    set_bg(sl, NAVY)

    # Orange top + right accent
    rect(sl, 0, 0, W, Pt(6), fill=ORANGE)
    rect(sl, W - Pt(6), 0, Pt(6), H, fill=ORANGE)

    # Chapter label
    ch = s.get('章番号', '01')
    txbox(sl, ML, Inches(2.0), CW, Pt(22),
          f'PART  {ch}', size=13, bold=True, color=MID_W)

    # Title
    title = s.get('タイトル', '')
    txbox(sl, ML, Inches(2.45), CW, Inches(1.8),
          title, size=38, bold=True, color=WHITE)

    # Lead
    lead = s.get('リード', '')
    if lead:
        txbox(sl, ML, Inches(4.35), CW, Inches(0.9),
              lead, size=17, color=DIM_W)

    add_notes(sl, s.get('スピーカーノート', ''))
    return sl


def slide_metric(prs, s):
    sl = prs.slides.add_slide(prs.slide_layouts[6])
    set_bg(sl, WHITE)

    rect(sl, 0, 0, W, Pt(4), fill=NAVY)
    rect(sl, 0, H - Pt(4), W, Pt(4), fill=ORANGE)

    lead = s.get('リード', '')
    txbox(sl, 0, Inches(1.4), W, Inches(0.7),
          lead, size=18, color=SUB, align=PP_ALIGN.CENTER)

    number = s.get('数字', '')
    txbox(sl, 0, Inches(2.1), W, Inches(2.4),
          number, size=80, bold=True, color=NAVY, align=PP_ALIGN.CENTER)

    desc = s.get('説明', '')
    txbox(sl, 0, Inches(4.7), W, Inches(0.7),
          desc, size=21, color=TEXT, align=PP_ALIGN.CENTER)

    source = s.get('出典', '')
    if source:
        txbox(sl, 0, Inches(5.5), W, Inches(0.5),
              f'出典：{source}', size=11, color=SUB, align=PP_ALIGN.CENTER)

    add_notes(sl, s.get('スピーカーノート', ''))
    return sl


def slide_quote(prs, s):
    sl = prs.slides.add_slide(prs.slide_layouts[6])
    set_bg(sl, NAVY_L)

    rect(sl, 0, 0, W, Pt(4), fill=NAVY)

    # Decorative quote mark
    txbox(sl, Inches(0.5), Inches(0.2), Inches(2), Inches(2.2),
          '“', size=90, color=NAVY_20)

    phrase = s.get('フレーズ', '')
    txbox(sl, Inches(1.5), Inches(1.9), Inches(10.3), Inches(3.0),
          phrase, size=32, bold=True, color=NAVY, align=PP_ALIGN.CENTER)

    add_notes(sl, s.get('スピーカーノート', ''))
    return sl


def slide_compare(prs, s):
    sl = prs.slides.add_slide(prs.slide_layouts[6])
    set_bg(sl, WHITE)

    rect(sl, 0, 0, W, Pt(4), fill=NAVY)
    content_y = header(sl, s.get('タイトル', ''))

    zukai = s.get('図解', {})
    L = zukai.get('左', {})
    R = zukai.get('右', {})

    col_w = (CW - Inches(0.35)) / 2
    avail_h = H - content_y - Inches(0.25)
    lx = ML
    rx = ML + col_w + Inches(0.35)

    # Left card (neutral)
    rect(sl, lx, content_y, col_w, avail_h, fill=GRAY_L, line=LINE, line_w_pt=0.75)
    # Right card (navy light)
    rect(sl, rx, content_y, col_w, avail_h, fill=NAVY_L, line=NAVY_20, line_w_pt=0.75)

    pad = Inches(0.2)
    iy = content_y + pad

    # Left label + title
    txbox(sl, lx + pad, iy, col_w - pad * 2, Pt(18),
          L.get('label', ''), size=11, bold=True, color=SUB)
    iy += Pt(19)
    txbox(sl, lx + pad, iy, col_w - pad * 2, Pt(24),
          L.get('title', ''), size=15, bold=True, color=TEXT)
    iy += Pt(26)
    for item in L.get('items', []):
        txbox(sl, lx + pad, iy, col_w - pad * 2, Pt(24),
              f'✓  {item}', size=13, color=TEXT)
        iy += Pt(24)

    iy = content_y + pad
    # Right label + title
    txbox(sl, rx + pad, iy, col_w - pad * 2, Pt(18),
          R.get('label', ''), size=11, bold=True, color=NAVY)
    iy += Pt(19)
    txbox(sl, rx + pad, iy, col_w - pad * 2, Pt(24),
          R.get('title', ''), size=15, bold=True, color=NAVY)
    iy += Pt(26)
    for item in R.get('items', []):
        txbox(sl, rx + pad, iy, col_w - pad * 2, Pt(24),
              f'✓  {item}', size=13, color=NAVY)
        iy += Pt(24)

    add_notes(sl, s.get('スピーカーノート', ''))
    return sl


def slide_flow(prs, s):
    sl = prs.slides.add_slide(prs.slide_layouts[6])
    set_bg(sl, WHITE)

    rect(sl, 0, 0, W, Pt(4), fill=NAVY)
    content_y = header(sl, s.get('タイトル', ''))

    zukai = s.get('図解', {})
    steps = zukai.get('ステップ', [])
    n = len(steps)
    if n == 0:
        return sl

    gap = Inches(0.2)
    step_w = (CW - gap * (n - 1)) / n
    step_h = H - content_y - Inches(0.25)
    y = content_y + Inches(0.05)

    for i, step in enumerate(steps):
        x = ML + (step_w + gap) * i

        rect(sl, x, y, step_w, step_h, fill=NAVY_L)

        # Number bubble (text approximation)
        txbox(sl, x, y + Inches(0.2), step_w, Pt(28),
              str(i + 1), size=18, bold=True, color=NAVY, align=PP_ALIGN.CENTER)

        txbox(sl, x + Inches(0.1), y + Inches(0.7), step_w - Inches(0.2), Pt(26),
              step.get('label', ''), size=14, bold=True, color=NAVY, align=PP_ALIGN.CENTER)

        txbox(sl, x + Inches(0.1), y + Inches(1.25), step_w - Inches(0.2), Inches(1.5),
              step.get('desc', ''), size=12, color=SUB, align=PP_ALIGN.CENTER)

        if i < n - 1:
            ax = x + step_w + gap * 0.1
            txbox(sl, ax, y + step_h / 2 - Pt(8), gap * 0.8, Pt(16),
                  '▶', size=9, color=NAVY_20, align=PP_ALIGN.CENTER)

    add_notes(sl, s.get('スピーカーノート', ''))
    return sl


def slide_key_message(prs, s):
    sl = prs.slides.add_slide(prs.slide_layouts[6])
    set_bg(sl, NAVY)

    rect(sl, 0, 0, W, Pt(6), fill=ORANGE)
    rect(sl, 0, H - Pt(6), W, Pt(6), fill=ORANGE)
    rect(sl, W - Pt(6), 0, Pt(6), H, fill=ORANGE)

    msg = s.get('メッセージ', '')
    txbox(sl, Inches(1.0), Inches(1.9), Inches(11.3), Inches(2.5),
          msg, size=36, bold=True, color=WHITE, align=PP_ALIGN.CENTER)

    rect(sl, W / 2 - Inches(1.0), Inches(4.5), Inches(2.0), Pt(3), fill=MID_W)

    sub = s.get('サブメッセージ', '')
    if sub:
        txbox(sl, Inches(1.0), Inches(4.85), Inches(11.3), Inches(1.2),
              sub, size=15, color=DIM_W, align=PP_ALIGN.CENTER)

    add_notes(sl, s.get('スピーカーノート', ''))
    return sl


# ─── dispatcher ──────────────────────────────────────────

def build_slide(prs, s):
    t = s.get('タイプ', '')
    if t == 'タイトル':
        return slide_cover(prs, s)
    if t == '章扉':
        return slide_section(prs, s)
    if t == '強調数字':
        return slide_metric(prs, s)
    if t == '決め台詞':
        return slide_quote(prs, s)
    if t == 'キーメッセージ':
        return slide_key_message(prs, s)
    if t == '図解':
        z_type = s.get('図解', {}).get('タイプ', '')
        if z_type == '比較':
            return slide_compare(prs, s)
        if z_type == 'プロセスフロー':
            return slide_flow(prs, s)
    # Default: content
    return slide_content(prs, s)


# ─── main ────────────────────────────────────────────────

def main():
    here = os.path.dirname(os.path.abspath(__file__))
    spec_path = os.path.join(here, 'slide-spec.json')
    out_path = os.path.join(here, '契約書AIチェックセミナー_v2.pptx')

    with open(spec_path, 'r', encoding='utf-8') as f:
        data = json.load(f)

    prs = Presentation()
    prs.slide_width = W
    prs.slide_height = H

    slides = data.get('スライド', [])
    for s in slides:
        build_slide(prs, s)
        print(f"  {s.get('番号', '?'):>2}  {s.get('タイプ', '?'):<10}  {s.get('タイトル', s.get('メッセージ', ''))[:40]}")

    prs.save(out_path)
    print(f'\n[OK] {len(slides)} slides -> {out_path}')


if __name__ == '__main__':
    main()
